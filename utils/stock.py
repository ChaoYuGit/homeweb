import os

import openpyxl

STOCK_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "database", "report")

SECTIONS = [
    ("sync",  "Sync Status",        "cron_daily.log"),
    ("macro", "Macro Data",         "01_Daily_Macro.txt"),
    ("mywatch", "Watching Stock",   "02_Daily_MyWatch.txt"),
    ("fund",  "StockInFund",        "03_Daily_StockInFund.txt"),
    ("etf",   "ETF观察",            "04_Weekly_ETF.xlsx"),
    ("fund_weekly", "Fund Weekly",  "05_Weekly_Fund.xlsx"),
]


def _read_file(filename):
    path = os.path.join(STOCK_DIR, filename)
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return f.read()
    return None


def get_stock_files():
    reports = []
    for sid, title, filename in SECTIONS:
        content = _read_file(filename)
        if content is not None:
            reports.append({"id": sid, "title": title, "filename": filename, "content": content})
    return reports


def get_section(section_id):
    for sid, title, filename in SECTIONS:
        if sid == section_id:
            content = _read_file(filename)
            if content is not None:
                return {"id": sid, "title": title, "filename": filename, "content": content}
            return None
    return None


def get_etf_data():
    path = os.path.join(STOCK_DIR, "04_Weekly_ETF.xlsx")
    if not os.path.exists(path):
        return None
    wb = openpyxl.load_workbook(path)
    ws = wb["ETF观察"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]

    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    table1_headers = [str(header_row[i]) if header_row[i] is not None else "" for i in range(2, 9)]
    table2_headers = [str(header_row[i]) if header_row[i] is not None else "" for i in range(19, 28)]
    table3_headers = [str(header_row[i]) if header_row[i] is not None else "" for i in range(37, 49)]
    table4_headers = [str(header_row[i]) if header_row[i] is not None else "" for i in range(85, 93)]

    def fmt(v):
        if v is None:
            return ""
        v = float(v)
        return f"{v:.2f}" if v == int(v) else f"{v:.2f}"

    def fmt_pct(v):
        if v is None:
            return ""
        return f"{float(v):.2f}%"

    table1 = []
    table2 = []
    table3 = []
    table4 = []
    for r in rows:
        table1.append({"code": r[0], "name": r[1], "c1": fmt_pct(r[2]), "c2": fmt_pct(r[3]), "c3": fmt_pct(r[4]), "c4": fmt_pct(r[5]), "c5": fmt_pct(r[6]), "c6": fmt_pct(r[7]), "c7": fmt_pct(r[8]), "d1": fmt(r[9]), "d2": fmt(r[10]), "d3": fmt(r[11]), "d4": fmt(r[12]), "d5": fmt(r[13]), "d6": fmt(r[14]), "d7": fmt(r[15])})
        table2.append({"code": r[0], "name": r[1], "c1": fmt_pct(r[19]), "c2": fmt_pct(r[20]), "c3": fmt_pct(r[21]), "c4": fmt_pct(r[22]), "c5": fmt_pct(r[23]), "c6": fmt_pct(r[24]), "c7": fmt_pct(r[25]), "c8": fmt_pct(r[26]), "c9": fmt_pct(r[27]), "d1": fmt(r[28]), "d2": fmt(r[29]), "d3": fmt(r[30]), "d4": fmt(r[31]), "d5": fmt(r[32]), "d6": fmt(r[33]), "d7": fmt(r[34]), "d8": fmt(r[35]), "d9": fmt(r[36])})
        t3 = {"code": r[0], "name": r[1]}
        for i in range(12):
            t3[f"c{i+1}"] = fmt_pct(r[37 + i])
            t3[f"d{i+1}"] = fmt(r[49 + i])
        table3.append(t3)
        t4 = {"code": r[0], "name": r[1]}
        for i in range(8):
            t4[f"c{i+1}"] = fmt_pct(r[85 + i])
        table4.append(t4)

    data_date = rows[0][18] if rows else ""
    return {"table1": table1, "table2": table2, "table3": table3, "table4": table4, "date": data_date,
            "table1_headers": table1_headers, "table2_headers": table2_headers,
            "table3_headers": table3_headers, "table4_headers": table4_headers}


def get_fund_weekly_data():
    path = os.path.join(STOCK_DIR, "05_Weekly_Fund.xlsx")
    if not os.path.exists(path):
        return None
    mtime = os.path.getmtime(path)
    from datetime import datetime
    date_str = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d")

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    col_headers = [str(header_row[i]) if header_row[i] is not None else "" for i in range(5, 14)]

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    watch_rows = [r for r in rows if r[3] is not None and int(r[3]) == 1]

    def fmt(v):
        if v is None:
            return ""
        v = float(v)
        return f"{v:.2f}" if v == int(v) else f"{v:.2f}"

    def fmt_pct(v):
        if v is None:
            return ""
        return f"{float(v)*100:.2f}%"

    table = []
    for r in watch_rows:
        row = {"code": r[0], "name": r[1], "avg": fmt_pct(r[4])}
        for i in range(9):
            row[f"c{i+1}"] = fmt_pct(r[5 + i])
            row[f"d{i+1}"] = fmt(r[14 + i])
        table.append(row)

    return {"date": date_str, "table": table, "headers": col_headers}
